#!/usr/bin/env -S uv run --script
# /// script
# requires-python = ">=3.10"
# dependencies = [
#   "python-pptx",
# ]
# ///
"""
render_slides.py — Render slides.pptx from outline.json + figs/.

Usage:
    uv run render_slides.py <paper-dir>

Expects:
    <paper-dir>/outline.json    # see references/slide_schema.md
    <paper-dir>/figs/           # extracted figures (png/jpg)

Writes:
    <paper-dir>/slides.pptx     # 16:9 deck

Output is purely .pptx (no PDF). Default target: 12 slides incl. title.
"""
import argparse, json, sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ─── theme ──────────────────────────────────────────────────────────
W, H = Inches(13.333), Inches(7.5)  # 16:9
TITLE_FONT = "Helvetica Neue"
BODY_FONT = "Helvetica Neue"
ACCENT = RGBColor(0x0E, 0x4D, 0x92)
SUBTLE = RGBColor(0x60, 0x60, 0x60)
TEXT = RGBColor(0x1A, 0x1A, 0x1A)
MAX_BULLETS = 6
MAX_BULLET_CHARS = 110


def _styled_paragraph(p, text, size=18, color=TEXT, bold=False, bullet=True):
    p.text = f"• {text}" if bullet else text
    p.font.name = BODY_FONT
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.space_after = Pt(6)


def add_title_slide(prs, title, authors, subtitle=""):
    blank = prs.slide_layouts[6]
    s = prs.slides.add_slide(blank)
    # title
    box = s.shapes.add_textbox(Inches(0.8), Inches(2.4), Inches(11.7), Inches(2.0))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title or "(untitled)"
    p.font.name = TITLE_FONT
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = ACCENT
    # author
    if authors:
        ab = s.shapes.add_textbox(Inches(0.8), Inches(4.6), Inches(11.7), Inches(0.6))
        ap_ = ab.text_frame.paragraphs[0]
        ap_.text = authors[:200]
        ap_.font.name = BODY_FONT
        ap_.font.size = Pt(16)
        ap_.font.color.rgb = SUBTLE
    if subtitle:
        sb = s.shapes.add_textbox(Inches(0.8), Inches(5.3), Inches(11.7), Inches(1.5))
        sp = sb.text_frame.paragraphs[0]
        sp.text = subtitle
        sp.font.name = BODY_FONT
        sp.font.size = Pt(14)
        sp.font.color.rgb = SUBTLE
    return s


def add_content_slide(prs, title, bullets, fig_path: Path | None = None):
    blank = prs.slide_layouts[6]
    s = prs.slides.add_slide(blank)

    # title bar
    tb = s.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12.0), Inches(0.8))
    tp = tb.text_frame.paragraphs[0]
    tp.text = title
    tp.font.name = TITLE_FONT
    tp.font.size = Pt(28)
    tp.font.bold = True
    tp.font.color.rgb = ACCENT

    # body
    has_fig = fig_path is not None and fig_path.exists()
    body_w = Inches(7.5) if has_fig else Inches(12.0)
    body = s.shapes.add_textbox(Inches(0.7), Inches(1.5), body_w, Inches(5.5))
    btf = body.text_frame
    btf.word_wrap = True

    truncated = [b[:MAX_BULLET_CHARS] for b in bullets[:MAX_BULLETS]]
    for i, b in enumerate(truncated):
        p = btf.paragraphs[0] if i == 0 else btf.add_paragraph()
        _styled_paragraph(p, b, size=18)

    if has_fig and fig_path.suffix.lower() in (".png", ".jpg", ".jpeg"):
        try:
            s.shapes.add_picture(
                str(fig_path), Inches(8.5), Inches(1.5),
                width=Inches(4.4),
            )
        except Exception as e:
            print(f"⚠ couldn't insert fig {fig_path.name}: {e}", file=sys.stderr)
    return s


def render(outline: dict, paper_dir: Path) -> Path:
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    add_title_slide(
        prs,
        outline.get("title", ""),
        outline.get("authors", ""),
        outline.get("subtitle", ""),
    )

    figs_dir = paper_dir / "figs"
    for sl in outline.get("slides", []):
        fig = None
        if sl.get("figure"):
            cand = figs_dir / Path(sl["figure"]).name
            if cand.exists():
                fig = cand
        add_content_slide(prs, sl.get("title", ""), sl.get("bullets", []), fig)

    out = paper_dir / "slides.pptx"
    prs.save(out)
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("paper_dir", help="dir containing outline.json + figs/")
    args = ap.parse_args()

    pd = Path(args.paper_dir).resolve()
    outline_path = pd / "outline.json"
    if not outline_path.exists():
        sys.exit(f"ERR: {outline_path} not found. Have the agent write it first.")

    outline = json.loads(outline_path.read_text())
    n_slides = 1 + len(outline.get("slides", []))
    out = render(outline, pd)
    print(f"✓ {out}  ({n_slides} slides)")


if __name__ == "__main__":
    main()
